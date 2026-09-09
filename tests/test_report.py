from datetime import date

import pandas as pd
from openpyxl import load_workbook

from report_pipeline.report import build_workbook
from report_pipeline.transform import ReportData


def _sample_data() -> ReportData:
    raw = pd.DataFrame(
        [
            {"date": "2026-09-09", "app_name": "앱A", "dau": 200, "installs": 20, "revenue": 2.0},
            {"date": "2026-09-09", "app_name": "앱B", "dau": 50, "installs": 5, "revenue": 0.5},
        ],
        columns=["date", "app_name", "dau", "installs", "revenue"],
    )
    summary_sheet = pd.DataFrame(
        [
            {"app_name": "앱A", "dau": 200, "installs": 20, "revenue": 2.0},
            {"app_name": "앱B", "dau": 50, "installs": 5, "revenue": 0.5},
        ]
    )
    return ReportData(sheets={"원본": raw, "요약": summary_sheet}, summary={"row_count": 2})


def test_build_workbook_uses_date_in_filename(tmp_path):
    path = build_workbook(_sample_data(), tmp_path, date(2026, 9, 9))
    assert path.name == "mobileindex_report_2026-09-09.xlsx"
    assert path.exists()


def test_build_workbook_creates_out_dir_if_missing(tmp_path):
    target = tmp_path / "nested" / "build"
    path = build_workbook(_sample_data(), target, date(2026, 9, 9))
    assert path.parent == target
    assert path.exists()


def test_workbook_sheets_and_header_formatting(tmp_path):
    path = build_workbook(_sample_data(), tmp_path, date(2026, 9, 9))
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["원본", "요약"]
    sheet = workbook["원본"]
    assert sheet["A1"].value == "date"
    assert sheet["A1"].font.bold is True
    assert sheet.freeze_panes == "A2"


def test_workbook_row_values_roundtrip(tmp_path):
    path = build_workbook(_sample_data(), tmp_path, date(2026, 9, 9))
    workbook = load_workbook(path)
    sheet = workbook["요약"]
    assert sheet["A2"].value == "앱A"
    assert sheet["B2"].value == 200
